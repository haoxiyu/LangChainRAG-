"""文档解析:把上传文件转成纯文本,供后续分块与向量化。

支持格式由 settings.allowed_extensions 控制。
"""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

# 文本文件按此顺序尝试解码,覆盖 UTF-8 与国内常见的 GBK 编码
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "gbk", "latin-1")


def _decode_text(raw: bytes) -> str:
    for enc in TEXT_ENCODINGS:
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            # 标注页码,便于回答时指出出处
            pages.append(f"[第{i}页]\n{text}")
    return "\n\n".join(pages)


def parse_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # 表格内容容易被忽略,但商品参数常在表格里,必须提取
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"[工作表: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def parse_csv(path: Path) -> str:
    raw = path.read_bytes()
    text = _decode_text(raw)
    rows: list[str] = []
    for row in csv.reader(io.StringIO(text)):
        cells = [c.strip() for c in row if c.strip()]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def parse_plain(path: Path) -> str:
    return _decode_text(path.read_bytes())


# 扩展名 → 解析函数
_PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".csv": parse_csv,
    ".txt": parse_plain,
    ".md": parse_plain,
}


def parse_file(path: Path) -> str:
    """按扩展名选择解析器,返回纯文本。不支持的格式抛 ValueError。"""
    ext = path.suffix.lower()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"不支持的文件格式: {ext or '(无扩展名)'}")
    text = parser(path)
    logger.info("解析完成 %s -> %d 字符", path.name, len(text))
    return text


def supported_extensions() -> set[str]:
    return set(_PARSERS)
